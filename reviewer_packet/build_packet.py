#!/usr/bin/env python3
"""
Assemble the independent-reviewer packet as a zip file.

The packet contains exactly three files:
    README_FOR_REVIEWER.md           - instructions (kept in this folder)
    human_adjudication_blinded.xlsx  - copied from data/screened/
    screening_rubric.md              - copied from analysis/

It deliberately excludes every file that shows the automated or AI second-
rater decisions (human_adjudication_worklist.csv, two_rater_comparison.csv,
three_way_comparison.csv, gemini_ratings.csv, and so on). The script checks
the finished zip against that list and refuses to produce one that leaks.

Usage:
    python3 reviewer_packet/build_packet.py [--date YYYY-MM-DD]

Output:
    reviewer_packet/reviewer_packet_<date>.zip   (git-ignored)
"""
import argparse
import datetime as dt
import shutil
import zipfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
HERE = ROOT / "reviewer_packet"
SOURCES = {
    "human_adjudication_blinded.xlsx": ROOT / "data/screened/human_adjudication_blinded.xlsx",
    "screening_rubric.md": ROOT / "analysis/screening_rubric.md",
}
README = HERE / "README_FOR_REVIEWER.md"

# Column names that would reveal the stratum or either rater's decision.
FORBIDDEN_COLUMNS = {"stratum", "review_type", "machine_ta", "machine_decision",
                     "machine_justification", "rater2", "rater2_reason", "concordant",
                     "rater1_first_screen", "rater2_claude", "rater3_gemini"}
FORBIDDEN_FILES = {"human_adjudication_worklist.csv", "two_rater_comparison.csv",
                   "three_way_comparison.csv", "gemini_ratings.csv", "cross_rater_stats.txt",
                   "final_human_worklist.csv", "human_verification_worklist.csv",
                   "human_adjudication_blinded.csv"}


# Words that belong to the author-side files, not to anything the reviewer sees.
# Checked in the Guide sheet and the Adjudication header row (not in abstracts,
# where words like "AI" occur legitimately).
FORBIDDEN_GUIDE_TERMS = ["rater", "worklist", "two_rater", "three_way", "gemini", "claude", "stratum",
                         "concordant", "concordance", "discordant", "spot-check", "spot check",
                         "disputed", "parse_fail", "machine_ta", "machine_decision", "machine decision",
                         "_incl_", "_excl", "AI answers", "analysis/"]
# Filename tokens that must not appear anywhere, including in data rows.
FORBIDDEN_DATA_TOKENS = ["worklist", "two_rater", "three_way", "gemini_ratings", "cross_rater"]


def check_workbook(path):
    wb = load_workbook(path)
    ws = wb["Adjudication"]
    headers = {c.value for c in ws[1]}
    leaked = headers & FORBIDDEN_COLUMNS
    if leaked:
        raise SystemExit(f"workbook leaks columns: {sorted(leaked)}")
    guide_text = " ".join(str(c.value) for row in wb["Guide"].iter_rows() for c in row if c.value)
    header_text = " ".join(str(h) for h in headers if h)
    hits = [t for t in FORBIDDEN_GUIDE_TERMS if t.lower() in (guide_text + " " + header_text).lower()]
    if hits:
        raise SystemExit(f"Guide sheet or headers contain author-side terms: {hits}")
    if "screening_rubric.md" not in guide_text:
        raise SystemExit("Guide sheet must point the reviewer to screening_rubric.md in the packet")
    data_text = " ".join(str(c.value) for row in ws.iter_rows(min_row=2) for c in row if c.value).lower()
    hits = [t for t in FORBIDDEN_DATA_TOKENS if t in data_text]
    if hits:
        raise SystemExit(f"Adjudication rows mention author-side files: {hits}")
    filled = [ws.cell(row=r, column=c).value for r in range(2, ws.max_row + 1)
              for c, name in enumerate(ws[1], 1) if name.value in
              ("your_decision", "your_reason_code", "your_notes")]
    if any(v not in (None, "") for v in filled):
        raise SystemExit("workbook answer columns are not blank; rebuild it first")
    return ws.max_row - 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date", default=dt.date.today().isoformat())
    args = ap.parse_args()

    for name, src in SOURCES.items():
        shutil.copyfile(src, HERE / name)
    n = check_workbook(HERE / "human_adjudication_blinded.xlsx")

    out = HERE / f"reviewer_packet_{args.date}.zip"
    members = [README.name, *SOURCES]
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name in members:
            z.write(HERE / name, arcname=f"reviewer_packet/{name}")

    with zipfile.ZipFile(out) as z:
        names = {Path(i).name for i in z.namelist()}
    bad = names & FORBIDDEN_FILES
    if bad:
        out.unlink()
        raise SystemExit(f"zip would leak files: {sorted(bad)}")
    print(f"packet: {out.relative_to(ROOT)}  ({len(members)} files, {n} records, blank answers verified)")


if __name__ == "__main__":
    main()

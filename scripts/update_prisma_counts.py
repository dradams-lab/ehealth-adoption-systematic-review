#!/usr/bin/env python3
"""
update_prisma_counts.py — Recompute PRISMA flow counts from the screening log.

Reads data/screened/screening_log_combined.csv plus
data/screened/duplicates_removed.csv, computes counts at each PRISMA
2020 stage, and updates prisma/prisma-counts.xlsx accordingly.

Decision-column conventions in screening_log_combined.csv:
    TA_Decision        = "Include" | "Exclude" | "" (unscreened)
    FullText_Decision  = "Include" | "Exclude" | "Not retrievable" | ""

PRISMA stage counts produced:
    Identification              = combined raw, pre-dedup    (= screening_log + duplicates_removed)
    Duplicates removed
    After deduplication         = unique records
    Excluded (T/A)              = TA_Decision == "Exclude"
    Sought for full-text        = TA_Decision == "Include"
    Not retrievable             = FullText_Decision == "Not retrievable"
    Full-text assessed          = sought - not_retrievable
    Excluded (full-text)        = FullText_Decision == "Exclude"
    Included in synthesis       = FullText_Decision == "Include"

Usage:
    python scripts/update_prisma_counts.py
"""
from __future__ import annotations
import sys
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
LOG = ROOT / "data" / "screened" / "screening_log_combined.csv"
DUP = ROOT / "data" / "screened" / "duplicates_removed.csv"
XLSX = ROOT / "prisma" / "prisma-counts.xlsx"

# Maps PRISMA-stage label (left col of the xlsx) -> the row's "n" cell to update.
# Adjust if you reorder rows in the xlsx.
ROW_LABEL_TO_UPDATE = {
    "TOTAL RECORDS IDENTIFIED (pre-deduplication)": "Identification",
    "Duplicate records removed": "Duplicates removed",
    "Records after deduplication": "After deduplication",
    "Records excluded — title/abstract screening": "Excluded T/A",
    "Records sought for full-text retrieval": "Sought full-text",
    "Full texts not retrievable": "Not retrievable",
    "Full texts assessed for eligibility": "Full-text assessed",
    "Full texts excluded after review": "Excluded full-text",
    "Studies included in qualitative synthesis": "Included",
}

def main() -> int:
    if not LOG.exists():
        sys.exit(f"Missing {LOG}. Run scripts/dedupe_and_combine.py first.")
    log = pd.read_csv(LOG, dtype=str).fillna("")
    dup = pd.read_csv(DUP, dtype=str).fillna("") if DUP.exists() else pd.DataFrame()

    counts = {
        "Identification":      len(log) + len(dup),
        "Duplicates removed":  len(dup),
        "After deduplication": len(log),
        "Excluded T/A":        int((log["TA_Decision"].str.lower() == "exclude").sum()),
        "Sought full-text":    int((log["TA_Decision"].str.lower() == "include").sum()),
        "Not retrievable":     int((log["FullText_Decision"].str.lower() == "not retrievable").sum()),
        "Excluded full-text":  int((log["FullText_Decision"].str.lower() == "exclude").sum()),
        "Included":            int((log["FullText_Decision"].str.lower() == "include").sum()),
    }
    counts["Full-text assessed"] = counts["Sought full-text"] - counts["Not retrievable"]

    print("Computed PRISMA stage counts:")
    for k, v in counts.items():
        print(f"  {k:25s} {v}")

    # Update xlsx in place — find each row by its first-column label
    if not XLSX.exists():
        print(f"\n(skip xlsx update — {XLSX} not present)")
        return 0
    wb = load_workbook(XLSX)
    ws = wb["PRISMA Flow Counts"]
    updates = 0
    for row in ws.iter_rows(min_row=1):
        label = (row[0].value or "").strip()
        if label in ROW_LABEL_TO_UPDATE:
            counts_key = ROW_LABEL_TO_UPDATE[label]
            new_value = counts[counts_key]
            ws.cell(row=row[0].row, column=2, value=new_value)
            updates += 1
    wb.save(XLSX)
    print(f"\nUpdated {updates} cells in {XLSX.relative_to(ROOT)}")
    return 0

if __name__ == "__main__":
    sys.exit(main())

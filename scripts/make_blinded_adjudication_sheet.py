#!/usr/bin/env python3
"""
Build a blinded human-adjudication sheet from the 126-record worklist.

The unblinded worklist (data/screened/human_adjudication_worklist.csv) shows
the primary screen's decision, the AI second rater's decision, the stratum,
and whether the record is disputed. Seeing those anchors the human rater, so
this script writes a sheet with all of them removed, adds the abstract from
the corpus, and shuffles the row order (fixed seed) so strata are not grouped.

Inputs:
    data/screened/human_adjudication_worklist.csv
    data/screened/corpus_combined.csv
Outputs (same rows, same order):
    data/screened/human_adjudication_blinded.xlsx  (preferred: dropdowns)
    data/screened/human_adjudication_blinded.csv   (UTF-8 with BOM, fallback)

Fill in ONE of the two files, not both. Join the completed sheet back to the
worklist on rec_id to unblind.
"""
import csv
import math
import random
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
WORKLIST = ROOT / "data/screened/human_adjudication_worklist.csv"
CORPUS = ROOT / "data/screened/corpus_combined.csv"
OUT = ROOT / "data/screened/human_adjudication_blinded.csv"
OUT_XLSX = ROOT / "data/screened/human_adjudication_blinded.xlsx"
SEED = 20260929

DECISIONS = ["INCLUDE", "EXCLUDE", "MAYBE"]
# E5 (duplicate / non-English) is omitted: it was never used in the screen.
REASON_CODES = [
    ("E1", "Stage 1", "Wrong technology focus: not about eHealth/EHR/HIE/health-IT/interoperability adoption."),
    ("E2", "Stage 1", "About eHealth, but not about adoption or implementation determinants."),
    ("E3", "Stage 1", "Out-of-scope setting or population."),
    ("E4", "Stage 1", "Ineligible publication type: editorial, letter, commentary, abstract without methods, protocol, poster."),
    ("E6", "Stage 1", "Published before 2015."),
    ("X1", "Stage 2", "Patient- or consumer-facing acceptance only; no organizational or system-level adoption lens."),
    ("X2", "Stage 2", "Single-condition digital-health intervention, not adoption of an eHealth system."),
    ("X3", "Stage 2", "EHR used only as a data source or delivery channel for a clinical outcome."),
    ("X4", "Stage 2", "Out-of-scope technology or setting."),
    ("X5", "Stage 2", "Non-research publication type."),
]

FIELDS = ["order", "rec_id", "title", "year", "venue", "doctype", "doi_link",
          "abstract", "your_decision", "your_reason_code", "your_notes"]


def split_abstract(text, limit):
    """Split text into chunks of at most `limit` characters, at sentence ends where possible."""
    chunks = []
    while len(text) > limit:
        cut = text.rfind(". ", int(limit * 0.6), limit)
        cut = cut + 1 if cut != -1 else text.rfind(" ", 0, limit)
        chunks.append(text[:cut].strip())
        text = text[cut:].strip()
    chunks.append(text)
    return chunks


def write_xlsx(out_rows):
    """Write the same rows as an Excel workbook with dropdowns and a guide sheet.

    Layout differs from the CSV for readability: the three answer columns sit
    on the left and stay frozen, and abstracts too long for one Excel cell
    (row height is capped at 409 pt) continue in abstract_continued columns.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    input_fill = PatternFill("solid", fgColor="FFFF00")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    top_wrap = Alignment(vertical="top", wrap_text=True)

    ABSTRACT_WIDTH, ABSTRACT_LIMIT = 100, 2700
    MARK = " [continues in next column]"
    chunked = [split_abstract(r["abstract"], ABSTRACT_LIMIT) for r in out_rows]
    n_extra = max(len(c) for c in chunked) - 1
    extra = ["abstract_continued"] + [f"abstract_continued_{i}" for i in range(2, n_extra + 1)]
    extra = extra[:n_extra]

    columns = (["order", "rec_id", "your_decision", "your_reason_code", "your_notes",
                "year", "doctype", "venue", "title", "abstract"] + extra + ["doi_link"])
    widths = {"order": 6, "rec_id": 8, "your_decision": 15, "your_reason_code": 19,
              "your_notes": 22, "year": 6, "doctype": 14, "venue": 16, "title": 28,
              "abstract": ABSTRACT_WIDTH, "doi_link": 30}
    widths.update({name: ABSTRACT_WIDTH for name in extra})
    inputs = {"your_decision", "your_reason_code", "your_notes"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudication"
    ws.append(columns)
    col = {}
    for i, name in enumerate(columns, 1):
        col[name] = get_column_letter(i)
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[col[name]].width = widths[name]
    ws.row_dimensions[1].height = 30

    for r, (row, chunks) in enumerate(zip(out_rows, chunked), 2):
        values = dict(row)
        for i, name in enumerate(["abstract"] + extra):
            text = chunks[i] if i < len(chunks) else ""
            values[name] = text + (MARK if i < len(chunks) - 1 else "")
        lines = 2
        for i, name in enumerate(columns, 1):
            value = values[name]
            if name in ("order", "year") and str(value).isdigit():
                value = int(value)
            cell = ws.cell(row=r, column=i, value=value or None)
            cell.font, cell.border, cell.alignment = font, border, top_wrap
            if name in inputs:
                cell.fill = input_fill
            if name == "doi_link" and value:
                cell.hyperlink, cell.font = value, link_font
            if name != "doi_link":
                # Arial 10 fits ~1.15 characters per width unit; assume 1.05, plus 8% for word wrap
                lines = max(lines, math.ceil(len(str(value or "")) / (widths[name] * 1.05) * 1.08))
        ws.row_dimensions[r].height = min(409, 13 * lines + 6)

    last = len(out_rows) + 1
    dv_dec = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Pick from the list",
                            error="Choose INCLUDE, EXCLUDE, or MAYBE.")
    dv_code = DataValidation(type="list",
                             formula1='"' + ",".join(c for c, _, _ in REASON_CODES) + '"',
                             allow_blank=True, showErrorMessage=True,
                             errorTitle="Pick from the list",
                             error="Choose one reason code (see the Guide sheet). Leave blank unless EXCLUDE.")
    ws.add_data_validation(dv_dec)
    ws.add_data_validation(dv_code)
    dv_dec.add(f"{col['your_decision']}2:{col['your_decision']}{last}")
    dv_code.add(f"{col['your_reason_code']}2:{col['your_reason_code']}{last}")

    ws.freeze_panes = f"{col['year']}2"  # header row and the answer columns stay visible
    ws.auto_filter.ref = f"A1:{col[columns[-1]]}{last}"

    # ---- Guide sheet ----
    g = wb.create_sheet("Guide")
    g.column_dimensions["A"].width = 34
    g.column_dimensions["B"].width = 14
    g.column_dimensions["C"].width = 100
    dec, code = col["your_decision"], col["your_reason_code"]
    rng_dec = f"Adjudication!${dec}$2:${dec}${last}"
    rng_code = f"Adjudication!${code}$2:${code}${last}"

    def put(r, a=None, b=None, c=None, style=font):
        for letter, v in (("A", a), ("B", b), ("C", c)):
            if v is not None:
                cell = g[f"{letter}{r}"]
                cell.value, cell.font, cell.alignment = v, style, top_wrap

    put(1, "Human adjudication: blinded sheet", style=Font(name="Arial", size=13, bold=True))
    put(3, "Progress", style=bold)
    put(4, "Records in the sheet", f"=COUNTA(Adjudication!$B$2:$B${last})")
    put(5, "Decisions entered", f"=COUNTA({rng_dec})")
    put(6, "Still to do", "=B4-B5")
    put(7, "EXCLUDE with no reason code", f'=COUNTIFS({rng_dec},"EXCLUDE",{rng_code},"")',
        "Should be 0 when you finish.")

    put(9, "How to fill it in", style=bold)
    steps = [
        "Fill in only the three yellow columns on the Adjudication sheet: your_decision, your_reason_code, your_notes.",
        "Decide on the title and abstract only (plus year, venue, and document type), not the full paper. If there is no abstract, use the title, venue, and document type.",
        "A long abstract continues in the abstract_continued columns; the first part ends with \"[continues in next column]\". Read all parts.",
        "your_decision: pick INCLUDE, EXCLUDE, or MAYBE from the dropdown. MAYBE means you cannot tell without the full text.",
        "your_reason_code: pick one code, only when your decision is EXCLUDE.",
        "your_notes: optional. A few words when a call is borderline.",
        "Scope test: is the record about adoption or implementation of eHealth systems (EHR/EMR/HIE/health-IT infrastructure, including interoperability) at the organizational or national/system level? The level refers to the system being adopted, not to whether the respondents are individual clinicians or staff (README condition 6).",
        "This workbook contains nothing that shows how the automated screen decided any record. Decide each record yourself, and please do not ask the author about a specific record until you have finished (README section 9).",
        "To see what is left, filter your_decision for blanks. Do not delete rows or edit the other columns. Save often. When you have finished, save under the same name with your initials added, for example human_adjudication_blinded_JD.xlsx (README section 10).",
    ]
    for i, text in enumerate(steps, 1):
        put(9 + i, f"Step {i}", text)
        g.merge_cells(start_row=9 + i, start_column=2, end_row=9 + i, end_column=3)
        g.row_dimensions[9 + i].height = 30

    r0 = 9 + len(steps) + 2
    put(r0, "Example of a filled-in row (made-up record, not in the sample)", style=bold)
    put(r0 + 1, "title", c="Patient satisfaction with a diabetes self-management smartphone app: a survey")
    put(r0 + 2, "your_decision", c="EXCLUDE")
    put(r0 + 3, "your_reason_code", c="X1")
    put(r0 + 4, "your_notes", c="Individual app acceptance; no organizational adoption lens")

    r1 = r0 + 6
    put(r1, "Reason codes", style=bold)
    put(r1 + 1, "Code", "Stage", "Meaning", style=bold)
    for i, (c_, stage, meaning) in enumerate(REASON_CODES, r1 + 2):
        put(i, c_, stage, meaning)
    put(r1 + 2 + len(REASON_CODES) + 1, "Source",
        c="Definitions: screening_rubric.md (in this folder) and README section 7. E5 (duplicate / non-English) is left out because the screen never used it.")

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT_XLSX)
    print(f"wrote {len(out_rows)} rows to {OUT_XLSX.relative_to(ROOT)}")


def main():
    corpus = {r["rec_id"]: r for r in csv.DictReader(open(CORPUS, encoding="utf-8"))}
    rows = list(csv.DictReader(open(WORKLIST, encoding="utf-8")))
    random.Random(SEED).shuffle(rows)

    out_rows = []
    with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for i, r in enumerate(rows, 1):
            c = corpus[r["rec_id"]]
            doi = c["doi"].strip()
            out_rows.append({
                "order": i,
                "rec_id": r["rec_id"],
                "title": c["title"],
                "year": c["year"],
                "venue": c["venue"],
                "doctype": c["doctype"],
                "doi_link": f"https://doi.org/{doi}" if doi else "",
                "abstract": c["abstract"].strip() or "[no abstract - screen on title, venue, and document type]",
                "your_decision": "",
                "your_reason_code": "",
                "your_notes": "",
            })
            w.writerow(out_rows[-1])
    print(f"wrote {len(rows)} rows to {OUT.relative_to(ROOT)}")
    write_xlsx(out_rows)


if __name__ == "__main__":
    main()

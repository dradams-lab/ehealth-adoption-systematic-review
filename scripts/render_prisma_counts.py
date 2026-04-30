#!/usr/bin/env python3
"""
render_prisma_counts.py — Print LaTeX-friendly PRISMA counts.

Reads prisma/prisma-counts.xlsx and prints two snippets:

  1. Figure 1 (TikZ) values — copy-paste into manuscript/figures/prisma_tikz.tex
  2. Table 2 Count column — copy-paste into manuscript/sections/prisma_flow.tex

Usage:
    python scripts/render_prisma_counts.py
"""
from __future__ import annotations
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "prisma" / "prisma-counts.xlsx"

ROW_LABELS_OF_INTEREST = [
    "TOTAL RECORDS IDENTIFIED (pre-deduplication)",
    "Duplicate records removed",
    "Records after deduplication",
    "Records excluded — title/abstract screening",
    "Records sought for full-text retrieval",
    "Full texts not retrievable",
    "Full texts assessed for eligibility",
    "Full texts excluded after review",
    "Studies included in qualitative synthesis",
    "Purposively selected national cases",
]

def main() -> int:
    if not XLSX.exists():
        sys.exit(f"Missing {XLSX}")

    wb = load_workbook(XLSX, data_only=True)
    ws = wb["PRISMA Flow Counts"]
    counts = {}
    for row in ws.iter_rows(min_row=1):
        label = (row[0].value or "").strip()
        if label in ROW_LABELS_OF_INTEREST:
            counts[label] = row[1].value

    g = lambda k: counts.get(k, "?")
    identified  = g("TOTAL RECORDS IDENTIFIED (pre-deduplication)")
    dup_removed = g("Duplicate records removed")
    after_dedup = g("Records after deduplication")
    excluded_ta = g("Records excluded — title/abstract screening")
    sought      = g("Records sought for full-text retrieval")
    not_retrieve = g("Full texts not retrievable")
    assessed    = g("Full texts assessed for eligibility")
    excluded_ft = g("Full texts excluded after review")
    included    = g("Studies included in qualitative synthesis")
    cases       = g("Purposively selected national cases")

    print("=" * 70)
    print("Figure 1 (TikZ) — copy values into manuscript/figures/prisma_tikz.tex")
    print("=" * 70)
    print(f"  Identification:      $n = {identified}$")
    print(f"  After dedup:         $n = {after_dedup}$")
    print(f"  Screened (T/A):      $n = {after_dedup}$")
    print(f"  Excluded (T/A):      $n = {excluded_ta}$")
    print(f"  Full-text assessed:  $n = {assessed}$  ({sought} sought; {not_retrieve} not retrievable)")
    print(f"  Full-text excluded:  $n = {excluded_ft}$")
    print(f"  Included:            $n = {included}$")
    print(f"  Cases:               $n = {cases}$")

    print("\n" + "=" * 70)
    print("Table 2 (PRISMA flow summary) — Count column values")
    print("=" * 70)
    rows = [
        ("Identification",        identified),
        ("Deduplication",         dup_removed),
        ("Screening",             after_dedup),
        ("Exclusion (T/A)",       excluded_ta),
        ("Eligibility (sought)",  sought),
        ("Not Retrievable",       not_retrieve),
        ("Full-text Assessed",    assessed),
        ("Full-text Exclusion",   excluded_ft),
        ("Included",              included),
        ("Case Sources",          cases),
    ]
    for label, val in rows:
        print(f"  {label:25s} {val}")

    return 0

if __name__ == "__main__":
    sys.exit(main())
